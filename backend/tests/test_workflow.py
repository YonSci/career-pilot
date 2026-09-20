from datetime import datetime, timezone, timedelta
from io import BytesIO
from zipfile import ZipFile
import json
from docx import Document
from pypdf import PdfReader
from career.db import Session, Record, put, read, now
from career.service import expired, fingerprint, run_scan
from career.notifications import notify
from career.config import settings

JOB = {
    "title": "Climate Data Scientist",
    "company": "Example Research",
    "location": "Remote",
    "description": "Build Python climate data science pipelines and geospatial models. Experience with hydrology is preferred.",
    "url": "https://example.org/jobs/123",
    "deadline": None,
}
CV = "Example Applicant\nBuilt Python data science pipelines for climate and hydrology projects.\nDeveloped geospatial analysis and remote sensing training materials."


def verified_profile(client):
    response = client.post("/api/profile/text", json={"text": CV})
    assert response.status_code == 200
    p = response.json()
    assert all(not f["verified"] for f in p["facts"])
    p["name"] = "Example Applicant"
    for f in p["facts"]:
        f["verified"] = True
    assert client.put("/api/profile", json=p).status_code == 200
    return p


def package():
    return {
        "cv": {
            "title": "Example Applicant",
            "paragraphs": [
                {"text": "Built Python climate pipelines.", "evidence_ids": ["F2"]}
            ],
        },
        "cover_letter": {
            "title": "Application",
            "paragraphs": [
                {
                    "text": "I would welcome the opportunity to contribute.",
                    "evidence_ids": [],
                }
            ],
        },
        "answers": [
            {
                "question": "Describe your experience",
                "answer": "Built Python pipelines.",
                "evidence_ids": ["F2"],
                "character_limit": 100,
            }
        ],
        "additional_documents": [],
        "checklist": ["Review all documents."],
        "missing_information": ["Confirm work authorization."],
    }


def test_auth_required(client):
    assert (
        client.get("/api/state", headers={"Authorization": "Bearer wrong"}).status_code
        == 401
    )
    assert client.get("/health").json()["status"] == "ok"


def test_import_verify_dedupe_match_and_persistence(client):
    j = client.post("/api/jobs", json=JOB).json()
    assert client.post(f"/api/jobs/{j['id']}/match").status_code == 409
    p = verified_profile(client)
    j2 = client.post(
        "/api/jobs", json={**JOB, "url": JOB["url"] + "?utm_source=linkedin"}
    ).json()
    assert j["id"] == j2["id"] and not j2["created_new"]
    result = client.post(f"/api/jobs/{j['id']}/match").json()
    assert result["match"]["mode"] == "keyword"
    assert result["match"]["score"] > 0
    client.put(f"/api/jobs/{j['id']}/decision", json={"status": "saved"})
    with Session() as db:
        stored = db.get(Record, j["id"])
        assert stored.data["status"] == "saved"
    client.put("/api/profile", json=p)
    assert client.get("/api/state").json()["jobs"][0]["match"] is not None
    p["facts"][0]["text"] += " (edited)"
    client.put("/api/profile", json=p)
    assert client.get("/api/state").json()["jobs"][0]["match"] is None


def test_upload_docx_and_reject_scanned_empty(client):
    d = Document()
    d.add_paragraph(CV)
    b = BytesIO()
    d.save(b)
    r = client.post(
        "/api/profile/upload",
        files={
            "file": (
                "CV.docx",
                b.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert r.status_code == 200 and r.json()["facts"]
    assert (
        client.post(
            "/api/profile/upload", files={"file": ("bad.exe", b"garbage")}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/profile/upload", files={"file": ("empty.txt", b"")}
        ).status_code
        == 422
    )


def test_deadlines_and_unsafe_urls(client):
    assert expired({"deadline": "2000-01-01"})
    assert not expired({"deadline": None})
    assert not expired({"deadline": "not a date"})
    assert not expired({"deadline": datetime.now(timezone.utc).date().isoformat()})
    assert expired(
        {"deadline": (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()}
    )
    assert (
        client.post("/api/jobs", json={**JOB, "url": "javascript:alert(1)"}).status_code
        == 422
    )
    assert (
        client.post(
            "/api/sources", json={"kind": "greenhouse", "value": "../../etc/passwd"}
        ).status_code
        == 422
    )


def test_approval_generation_and_export(client, monkeypatch):
    verified_profile(client)
    j = client.post("/api/jobs", json=JOB).json()
    assert client.post(f"/api/jobs/{j['id']}/prepare").status_code == 409
    monkeypatch.setattr(settings, "openai_api_key", "fake-test-key")
    calls = []

    def write(*args):
        calls.append(True)
        return package()

    monkeypatch.setattr("career.service.write_package", write)
    assert client.post(f"/api/jobs/{j['id']}/prepare").status_code == 200
    app = client.get("/api/state").json()["applications"][0]
    assert (
        app["status"] == "review"
        and app["profile_snapshot"]["name"] == "Example Applicant"
    )
    client.post(f"/api/jobs/{j['id']}/prepare")
    assert len(calls) == 1
    response = client.get(f"/api/applications/{app['id']}/download")
    assert response.status_code == 200
    z = ZipFile(BytesIO(response.content))
    assert {
        "CV.docx",
        "CV.pdf",
        "Cover_letter.docx",
        "Screening_answers.docx",
        "Checklist.txt",
    }.issubset(z.namelist())
    assert "Python" in "\n".join(
        p.text for p in Document(BytesIO(z.read("CV.docx"))).paragraphs
    )
    assert "Python" in PdfReader(BytesIO(z.read("CV.pdf"))).pages[0].extract_text()
    invalid = package()
    invalid["answers"][0]["character_limit"] = 2
    assert (
        client.put(
            f"/api/applications/{app['id']}",
            json={"package": invalid, "status": "ready"},
        ).status_code
        == 422
    )
    assert (
        client.put(
            f"/api/applications/{app['id']}",
            json={"package": package(), "status": "ready"},
        ).json()["status"]
        == "ready"
    )


def test_failed_preparation_is_visible_and_retryable(client, monkeypatch):
    verified_profile(client)
    j = client.post("/api/jobs", json=JOB).json()
    monkeypatch.setattr(settings, "openai_api_key", "fake-test-key")

    def failed(*a):
        raise ValueError("Missing required information.")

    monkeypatch.setattr("career.service.write_package", failed)
    client.post(f"/api/jobs/{j['id']}/prepare")
    app = client.get("/api/state").json()["applications"][0]
    assert app["status"] == "failed" and "Missing" in app["error"]


def test_scans_and_notification_duplicate_suppression(client, monkeypatch):
    verified_profile(client)
    client.post("/api/sources", json={"kind": "greenhouse", "value": "example"})
    monkeypatch.setattr("career.service.collect", lambda *a: ([JOB], []))
    assert client.post("/api/scan").status_code == 200
    data = client.get("/api/state").json()
    assert len(data["jobs"]) == 1 and data["runs"][0]["status"] == "completed"
    client.post("/api/scan")
    assert len(client.get("/api/state").json()["jobs"]) == 1
    monkeypatch.setattr("career.notifications.availability", lambda db=None: {"telegram": True})
    sends = []
    monkeypatch.setattr("career.notifications.send_alert", lambda *a, **k: sends.append(a))
    with Session() as db:
        notify(db, "testjob", JOB, {"score": 80}, ["telegram"])
        notify(db, "testjob", JOB, {"score": 80}, ["telegram"])
    assert len(sends) == 1


def test_ambiguous_delivery_is_never_automatically_resent(client, monkeypatch):
    monkeypatch.setattr("career.notifications.availability", lambda db=None: {"email": True})
    calls = []

    def timeout(*a, **k):
        calls.append(a)
        raise TimeoutError()

    monkeypatch.setattr("career.notifications.send_alert", timeout)
    with Session() as db:
        first = notify(db, "ambiguous", JOB, {"score": 80}, ["email"])
        second = notify(db, "ambiguous", JOB, {"score": 80}, ["email"])
    assert (
        first["email"] == "delivery_unknown"
        and second["email"] == "already_attempted"
        and len(calls) == 1
    )


def test_telegram_webhook_rejects_wrong_identity(client, monkeypatch):
    monkeypatch.setattr(settings, "telegram_webhook_secret", "secret")
    monkeypatch.setattr(settings, "telegram_chat_id", "123")
    assert client.post("/api/telegram/webhook", json={}).status_code == 401
    payload = {
        "callback_query": {
            "from": {"id": 999},
            "message": {"chat": {"id": 123}},
            "data": "save:test",
        }
    }
    assert (
        client.post(
            "/api/telegram/webhook",
            json=payload,
            headers={"X-Telegram-Bot-Api-Secret-Token": "secret"},
        ).status_code
        == 403
    )
